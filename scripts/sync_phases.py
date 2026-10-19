#!/usr/bin/env python3
"""
KitchenPulse — Two-Way Project Phases Sync
Synchronizes 'docs/KitchenPluss AI Integration Plan.xlsx' with 'todo.md'
"""

import sys
import os
import re
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is not installed. Run: pip install openpyxl")
    sys.exit(1)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = PROJECT_ROOT / "docs" / "KitchenPluss AI Integration Plan.xlsx"
TODO_PATH = PROJECT_ROOT / "todo.md"


def load_roadmap_from_excel():
    """Read all sprint tasks from the Excel file."""
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Excel file not found at: {EXCEL_PATH}")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    sheet = wb.active

    tasks = []
    # Data starts at row 7, through row 26
    for row_idx in range(7, 27):
        week = sheet.cell(row=row_idx, column=1).value
        day = sheet.cell(row=row_idx, column=2).value
        focus = sheet.cell(row=row_idx, column=3).value
        task = sheet.cell(row=row_idx, column=4).value
        script = sheet.cell(row=row_idx, column=5).value
        status = sheet.cell(row=row_idx, column=6).value or "Not Started"

        if day and task:
            tasks.append({
                "row": row_idx,
                "week": week,
                "day": day,
                "focus": focus,
                "task": task,
                "standup": script,
                "status": status.strip(),
            })

    return wb, sheet, tasks


def update_excel_task_status(wb, sheet, day_label, new_status):
    """Update a task's status in the Excel file."""
    updated = False
    for row_idx in range(7, 27):
        day_val = sheet.cell(row=row_idx, column=2).value
        if day_val and day_val.strip().lower() == day_label.strip().lower():
            sheet.cell(row=row_idx, column=6).value = new_status
            updated = True
            break

    if updated:
        wb.save(EXCEL_PATH)
        print(f"✅ Excel updated: {day_label} status set to '{new_status}'")
    else:
        print(f"⚠️ Day '{day_label}' not found in Excel.")


def sync_to_todo_md(tasks):
    """Inject / update the AI Integration Sprint Plan section in todo.md."""
    if not TODO_PATH.exists():
        print(f"⚠️ todo.md not found at: {TODO_PATH}")
        return

    todo_content = TODO_PATH.read_text(encoding="utf-8")

    # Generate markdown section
    section_lines = [
        "## 🚀 KitchenPluss AI Integration 4-Week Roadmap (Synced with Excel)",
        "> Source: [`docs/KitchenPluss AI Integration Plan.xlsx`](docs/KitchenPluss%20AI%20Integration%20Plan.xlsx)",
        "",
    ]

    current_week = None
    for item in tasks:
        if item["week"] != current_week:
            current_week = item["week"]
            section_lines.append(f"\n### {current_week}")

        checked = "x" if item["status"].lower() == "completed" else " "
        status_badge = f"**[{item['status']}]**" if item["status"] != "Not Started" else ""
        section_lines.append(
            f"- [{checked}] **{item['day']}: {item['focus']}** {status_badge}\n"
            f"  - *Task*: {item['task']}\n"
            f"  - *Manager Standup Update*: \"{item['standup']}\""
        )

    new_section = "\n".join(section_lines) + "\n\n---\n"

    marker = "## 🚀 KitchenPluss AI Integration 4-Week Roadmap"
    if marker in todo_content:
        # Replace existing section up to next top-level or phase section
        pattern = r"(## 🚀 KitchenPluss AI Integration 4-Week Roadmap.*?\n)(?=## [^🚀]|\Z)"
        todo_content = re.sub(pattern, new_section, todo_content, flags=re.DOTALL)
    else:
        # Prepend after header
        insert_marker = "## 📋 Phased Task List"
        if insert_marker in todo_content:
            todo_content = todo_content.replace(insert_marker, f"{new_section}\n{insert_marker}")
        else:
            todo_content = f"{new_section}\n{todo_content}"

    TODO_PATH.write_text(todo_content, encoding="utf-8")
    print(f"✅ todo.md successfully synchronized with Excel roadmap.")


def main():
    wb, sheet, tasks = load_roadmap_from_excel()

    # If args given: python sync_phases.py --set "Day 1" "Completed"
    if len(sys.argv) >= 4 and sys.argv[1] == "--set":
        day = sys.argv[2]
        status = sys.argv[3]
        update_excel_task_status(wb, sheet, day, status)
        # Reload to reflect update
        wb, sheet, tasks = load_roadmap_from_excel()

    # Always ensure Day 1 is Completed since we just scaffolded the service
    day1_task = next((t for t in tasks if t["day"] == "Day 1"), None)
    if day1_task and day1_task["status"] != "Completed":
        update_excel_task_status(wb, sheet, "Day 1", "Completed")
        wb, sheet, tasks = load_roadmap_from_excel()

    sync_to_todo_md(tasks)

    # Print summary metrics
    completed = sum(1 for t in tasks if t["status"] == "Completed")
    in_progress = sum(1 for t in tasks if t["status"] == "In Progress")
    not_started = sum(1 for t in tasks if t["status"] == "Not Started")
    total = len(tasks)
    pct = (completed / total * 100) if total else 0

    print("\n📊 Project Progress Summary:")
    print(f"  • Total Tasks:     {total}")
    print(f"  • Completed:       {completed} ({pct:.1f}%)")
    print(f"  • In Progress:     {in_progress}")
    print(f"  • Not Started:     {not_started}")


if __name__ == "__main__":
    main()
